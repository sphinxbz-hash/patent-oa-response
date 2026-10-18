"""Build an editable, source-preserving response package without model calls."""

import csv
import html
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .analysis import analyze_case, document_text
from .claims import diff_tokens, parse_claims

DEMO_NOTICE = 'Fictional demonstration / not a real patent case'


def _xml_text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value or ''))


def _paragraphs(value):
    return [part.strip() for part in re.split(r'\n\s*\n|\r?\n', str(value or '')) if part.strip()]


def _source_label(source):
    if not source:
        return 'No usable source selected'
    label = f'{source.get("id", "")} | {source.get("country", "")} | {source.get("date") or "Date unavailable"}'
    if source.get('url'):
        label += ' | ' + source['url']
    return label


def _new_document(title):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(1)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.font.size = Pt(11 if name == 'Normal' else 19 if name == 'Title' else 13 if name == 'Heading 1' else 12)
        style.paragraph_format.space_after = Pt(8)
        if name != 'Normal':
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.keep_with_next = True
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.widow_control = True
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'en-US')
    normal.element.get_or_add_rPr().append(language)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.title = title
    document.add_paragraph(title, style='Title')
    return document


class _DualDocument:
    """One content stream keeps the Word document and browser preview aligned."""

    def __init__(self, title, is_demo=False):
        self.title = title
        self.word = _new_document(title)
        self.body = ['<h1>' + html.escape(title) + '</h1>']
        if is_demo:
            self.paragraph(DEMO_NOTICE)

    def paragraph(self, text):
        text = _xml_text(text)
        self.word.add_paragraph(text)
        self.body.append('<p>' + html.escape(text).replace('\n', '<br>') + '</p>')

    def heading(self, text):
        self.word.add_heading(_xml_text(text), level=1)
        self.body.append('<h2>' + html.escape(text) + '</h2>')

    def runs(self, prefix, runs):
        paragraph = self.word.add_paragraph()
        paragraph.add_run(_xml_text(prefix))
        markup = [html.escape(prefix)]
        for kind, text in runs:
            text = _xml_text(text)
            run = paragraph.add_run(text)
            escaped = html.escape(text).replace('\n', '<br>')
            if kind == 'insert':
                run.underline = True
                markup.append('<ins>' + escaped + '</ins>')
            elif kind == 'delete':
                run.font.strike = True
                markup.append('<del>' + escaped + '</del>')
            else:
                markup.append(escaped)
        self.body.append('<p class="claim">' + ''.join(markup) + '</p>')

    def save(self, folder, stem):
        self.word.save(folder / (stem + '.docx'))
        page = ('<!doctype html><html lang="en"><head><meta charset="utf-8">'
                '<meta name="viewport" content="width=device-width,initial-scale=1">'
                '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'">'
                '<title>' + html.escape(self.title) + '</title><style>'
                'body{max-width:800px;margin:40px auto;padding:0 24px;font:18px/1.55 Georgia,serif;color:#111;background:#fff;overflow-wrap:anywhere}'
                'h1{font-size:28px;line-height:1.25}h2{font-size:21px;margin-top:32px}'
                'p{white-space:normal}.claim{margin:22px 0}ins{text-decoration:underline}del{text-decoration:line-through}'
                '@media print{body{margin:0;max-width:none;font-size:11pt}h1,h2{break-after:avoid}}'
                '</style></head><body>' + ''.join(self.body) + '</body></html>')
        (folder / (stem + '.html')).write_text(page, encoding='utf-8')


def _claim_map(text):
    return {claim['number']: claim['text'] for claim in parse_claims(text)}


def _claims_content(case, analysis):
    settings, sources = case.get('settings') or {}, analysis['sources']
    override = str(settings.get('claims_override') or '').strip()
    text = override or document_text(sources['claims'])
    try:
        proposed = _claim_map(text)
    except ValueError:
        proposed = {}
        analysis['gaps'].append('Unable to parse proposed claims as one numbered claim set. Check duplicate claim numbers and version boundaries.')
    try:
        before = _claim_map(document_text(sources['pending'])) if analysis['baseline_usable'] else {}
    except ValueError:
        before = {}
        analysis['gaps'].append('Unable to parse target pending claims as one numbered claim set. No target redline is produced.')
    if text and not proposed:
        analysis['gaps'].append('The selected claim text has no recognised numbered claims; no claim listing was fabricated.')
    if not proposed:
        analysis['gaps'].append('No proposed numbered English claims are available.')
    if analysis['baseline_usable'] and not before:
        analysis['baseline_usable'] = False
        analysis['gaps'].append('The target pending text has no recognised numbered claims; no target redline was produced.')
    return proposed, before, bool(override)


def _write_response(case, analysis, folder, has_claims=False):
    output = _DualDocument('Response to the Office Action', case.get('is_demo'))
    output.paragraph(f'Target jurisdiction: {case.get("target_country") or "BR"}. '
                     f'Target application: {case.get("target_application") or "Not supplied"}.')
    output.paragraph('Draft response for the target application. ' + ('Please consider the accompanying amended claims and following observations.' if has_claims else 'Proposed claims are unavailable. The following source observations and evidence gaps form an incomplete working draft.'))
    sources = analysis['sources']
    output.heading('Target Office Action')
    output.paragraph(_source_label(sources['oa']))
    if analysis['issues']:
        output.paragraph('The following headings were detected by explicit text rules. They do not establish the merits of an objection.')
        for issue in analysis['issues']:
            page = str(issue['page']) if issue.get('page') is not None else 'unavailable'
            output.paragraph(f'{issue["title"]}. Source page: {page}. '
                             f'Citations: {", ".join(issue["citations"]) or "None detected"}.')
    else:
        output.paragraph('No supported target objection heading has been extracted. The target Office Action must be checked.')
    override = str((case.get('settings') or {}).get('response_override') or '').strip()
    if override:
        output.heading('User authored response')
        output.paragraph('The following text was supplied by the user and has not been independently verified.')
        for paragraph in _paragraphs(override):
            output.paragraph(paragraph)
    elif sources['response']:
        output.heading('Applicant argument adapted from the selected family response')
        output.paragraph(_source_label(sources['response']))
        output.paragraph('Known source application identifiers are adapted to the target application. Technical arguments are retained '
                         'from the applicant submission. The adaptation log and original text accompany this package.')
        for paragraph in _paragraphs(analysis['adaptation']['text']):
            output.paragraph(paragraph)
    else:
        output.heading('Argument required')
        output.paragraph('No source backed argument is available. Import an English family response or provide user authored text. '
                         'No substantive patentability argument has been generated.')
    output.heading('Proposed amendments')
    if has_claims:
        output.paragraph('The accompanying clean claims contain the selected English claim text or the user authored replacement. '
                         'The marked claims identify the available comparison baseline. The support table records citations to the target original filing.')
    else:
        output.paragraph('No proposed numbered English claims are available. The claim documents record the missing source; no amendment is asserted.')
    output.heading('Appendix of unresolved target specific items')
    for gap in analysis['gaps']:
        output.paragraph(gap)
    output.heading('Deadline and fee status')
    deadline = analysis['deadline']
    output.paragraph(deadline['explanation'])
    if deadline['nominal_date']:
        output.paragraph('Nominal date: ' + deadline['nominal_date'] + '. Rule source: ' + deadline['rule_source_url'])
    output.paragraph(deadline['fee']['explanation'])
    output.save(folder, '03-意见陈述书')


def _write_claims(case, analysis, folder, proposed, before, user_authored):
    marked = _DualDocument('Marked Claims', case.get('is_demo'))
    clean = _DualDocument('Clean Claims', case.get('is_demo'))
    for output in (marked, clean):
        output.paragraph(f'Target application: {case.get("target_application") or "Not supplied"}.')
        output.paragraph('User authored claim text.' if user_authored else 'Claim text reused from the selected English family claim set.')
        output.paragraph('Source: ' + _source_label(analysis['sources']['claims']))
    clean.paragraph('This is the proposed claim text without comparison marks. Target original support and entry remain to be verified.')
    has_baseline = bool(analysis['baseline_usable'] and before)
    is_us = str(case.get('target_country') or '').upper() == 'US'
    if has_baseline:
        marked.paragraph('Comparison baseline: ' + _source_label(analysis['sources']['pending']))
        marked.paragraph('Insertions are underlined and deletions are struck through. The baseline is the selected target pending version.')
        if is_us:
            marked.paragraph('US listing adapter. Confirm that the selected baseline is the immediately prior entered listing. '
                             'New claims are shown clean and canceled claims by number only. Unknown historical statuses are identified for verification.')
            if not analysis['baseline_latest']:
                marked.paragraph('The selected baseline is older than another supplied pending version. This comparison is not a verified US amendment listing.')
    else:
        marked.paragraph('No target redline is available because a usable English target pending baseline is missing. '
                         'The following is reference claim text only. It does not show amendments to the target application.')
    if not proposed:
        marked.paragraph('No numbered proposed claims are available. Import the complete English claim set or supply a numbered replacement.')
        clean.paragraph('No numbered proposed claims are available. No replacement claim language has been generated.')
    else:
        for number, text in sorted(proposed.items()):
            clean.paragraph(f'{number}. {text}')
        for number in sorted(set(before) | set(proposed)) if has_baseline else sorted(proposed):
            old, new = before.get(number), proposed.get(number)
            if not has_baseline:
                marked.paragraph(f'{number}. {new}')
            elif is_us and new is None:
                marked.paragraph(f'{number}. (Canceled)')
            elif is_us and old is None:
                marked.paragraph(f'{number}. (New) {new}')
            elif is_us and old == new:
                statuses = (analysis['sources']['pending'] or {}).get('claim_statuses') or {}
                status = statuses.get(str(number)) or statuses.get(number)
                if status not in ('Original', 'Previously presented', 'Withdrawn'):
                    status = 'Status unverified'
                marked.paragraph(f'{number}. ({status}) {new}')
            else:
                prefix = f'{number}. (Currently amended) ' if is_us else f'{number}. '
                marked.runs(prefix, diff_tokens(old or '', new or ''))
    marked.save(folder, '04-权利要求修改标记页')
    clean.save(folder, '05-修改后权利要求清洁稿')


def _md(value):
    return str(value if value is not None else '').replace('|', '\\|').replace('\r', '').replace('\n', '<br>')


def _support(analysis, proposed, before):
    rows = analysis['support']
    known_claims = {str(row['claim']) for row in rows}
    # Missing or malformed proposed text is not an instruction to cancel claims.
    for number in sorted(set(proposed) | set(before)) if proposed else []:
        if before.get(number) == proposed.get(number) or str(number) in known_claims:
            continue
        change = ('Proposed claim; target baseline unavailable' if not before else 'Canceled claim' if number not in proposed
                  else 'Added claim' if number not in before else 'Changed wording')
        rows.append(dict(claim=str(number), change=change, basis='', source_id='', page='', status='unverified',
                         legal_support='unverified', note='Provide a basis citation to the target original filing.'))
    missing = list(dict.fromkeys(str(row['claim']) for row in rows if row['status'] != 'source_located'))
    if missing:
        analysis['gaps'].append('Target original support remains unverified for claims ' + ', '.join(missing) + '.')
    return rows


def _write_support(analysis, folder, is_demo=False):
    columns = ['claim', 'change', 'basis', 'source_id', 'page', 'status', 'legal_support', 'note']
    rows = analysis['support']
    lines = ['# Amendment support in the target original', '',
             'Source located means that the supplied quotation appears on the cited target original page. '
             'It does not establish legal sufficiency, added matter compliance or claim scope support. '
             'Foreign documents are not verified target original support.', '',
             '| ' + ' | '.join(columns) + ' |', '| ' + ' | '.join('---' for _ in columns) + ' |']
    lines += ['| ' + ' | '.join(_md(row.get(key, '')) for key in columns) + ' |' for row in rows]
    if not rows:
        lines += ['', 'No amendment support rows are available.']
    (folder / '06-修改依据表.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    with (folder / '06-修改依据表.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        csv_columns = columns + ['case_notice'] if is_demo else columns
        writer = csv.DictWriter(handle, fieldnames=csv_columns, extrasaction='ignore')
        writer.writeheader()
        for row in rows or ([{}] if is_demo else []):
            # Keep spreadsheet viewers from executing user supplied formula text.
            safe = {key: ("'" + str(value) if str(value).lstrip().startswith(('=', '+', '-', '@')) else value)
                    for key, value in row.items()}
            if is_demo:
                safe['case_notice'] = DEMO_NOTICE
            writer.writerow(safe)


def _source_link(source):
    url = str((source or {}).get('url') or '')
    if url.startswith(('https://', 'http://')):
        return '[来源文件](' + url.replace('(', '%28').replace(')', '%29').replace(' ', '%20') + ')'
    return '来源链接未提供'


def _quote_text(text):
    fence = '`' * max(3, 1 + max((len(match.group()) for match in re.finditer(r'`+', text)), default=0))
    return [fence + 'text', text, fence, '']


def _write_history(case, analysis, folder):
    documents = case.get('documents') or []
    timestamps = [str(doc['retrieved_at']) for doc in documents if doc.get('retrieved_at')]
    retrieved = case.get('researched_at') or (max(timestamps) if timestamps else '未记录')
    saved = sum(bool(doc.get('local_path')) for doc in documents)
    readable = sum(bool(document_text(doc)) for doc in documents)
    lines = ['# 同族审查历史与复用线索', '']
    if case.get('is_demo'):
        lines += [DEMO_NOTICE, '']
    lines += ['输入检索号：' + str(case.get('input_number') or ''),
              '目标申请号：' + str(case.get('target_application') or '尚未填写'),
              '最近检索记录时间：' + str(retrieved), '',
              f'当前记录 {len(documents)} 份文件，其中 {saved} 份有本地原件路径，{readable} 份有可读取文本。'
              '案卷目录完整性未核实；本报告仅覆盖已识别材料，缺失记录不能用于推断文件不存在。', '',
              '申请人的论证与审查机关记录分别列示。外国授权状态本身不能证明某项论证获得采纳。', '',
              '## 同族成员', '']
    for row in case.get('family') or []:
        lines.append('- ' + ' | '.join(str(row.get(key) or 'Unknown') for key in ('country', 'application', 'publication', 'status', 'status_source', 'url')))
    if not case.get('family'):
        lines.append('尚无可用的同族成员证据。')
    lines += ['', '## 实际选用证据', '']
    for role, source in analysis['sources'].items():
        lines.append('- ' + role + ': ' + _source_label(source) + ' | ' + _source_link(source))
    lines += ['', '## 可复用论证清单', '',
              '下列内容是申请人原有论证的可追溯片段。对应关系来自明示引证或词语匹配，只是复用候选，不是法律适用性或成功率判断。', '']
    response = analysis['sources']['response']
    if response:
        for index, paragraph in enumerate(_paragraphs(document_text(response)), 1):
            normalized = re.sub(r'\s+', ' ', paragraph).strip()
            pages = [str(page['number']) for page in response.get('pages') or [] if page.get('number') is not None
                     and normalized in re.sub(r'\s+', ' ', str(page.get('text') or ''))]
            locator = '第 ' + '、'.join(pages) + ' 页' if pages else '页码未定位'
            matches = []
            for strategy in analysis['strategy']:
                evidence = strategy.get('evidence') or {}
                compact = re.sub(r'\s+', '', paragraph).upper()
                citation_match = any(citation in compact for citation in evidence.get('shared_citations') or [])
                matched_terms = [term for term in evidence.get('shared_terms') or []
                                 if re.search(r'\b' + re.escape(term) + r'\b', paragraph, re.I)]
                if strategy['status'] == 'reuse_argument' and (citation_match or len(matched_terms) >= 2):
                    matches.append(strategy['issue'])
            lines += [f'### 论证片段 {index}', '',
                      f'来源：{response.get("id", "")}，{locator}，{_source_link(response)}。',
                      '对应目标问题：' + ('、'.join(dict.fromkeys(matches)) + '（规则匹配候选）' if matches else '尚未定位') + '。', '']
            lines += _quote_text(paragraph)
    else:
        lines += ['尚未取得可用的同族英文答复，暂无法列出有来源的论证片段。', '']
    lines += ['## 目标问题与复用候选', '']
    status_labels = {'reuse_argument': '复用论证候选', 'reuse_amendment': '复用修改候选', 'separate_response': '需单独组织答复'}
    for row in analysis['strategy']:
        page = '第 ' + str(row['page']) + ' 页' if row.get('page') is not None else '页码未定位'
        lines.append('- ' + row['issue'] + ' | ' + status_labels.get(row['status'], row['status']) +
                     ' | 目标意见 ' + str(row['source_id']) + ' ' + page + ' | ' + _source_link(analysis['sources']['oa']))
    if not analysis['strategy']:
        lines.append('尚无可用的目标审查问题与论证映射。')
    lines += ['', '## 答复身份适配日志', '']
    if analysis['adaptation']['log']:
        for entry in analysis['adaptation']['log']:
            lines.append('- ' + entry['old'] + ' -> ' + entry['new'] + '; ' + str(entry['occurrences']) +
                         ' exact occurrence(s). ' + entry['explanation'])
    else:
        lines.append('未替换任何已知同族标识。用户自写文本不会被自动改动。')
    lines += ['', '## 文件时间线', '']
    for doc in sorted(case.get('documents') or [], key=lambda d: (str(d.get('date') or ''), str(d.get('id') or ''))):
        voice = '申请人提交' if doc.get('role') == 'response' else '审查机关意见记录' if doc.get('role') == 'oa' else '文件记录'
        lines.append('- ' + _source_label(doc) + ' | ' + voice + ' | ' + str(doc.get('role') or 'unknown') +
                     ' | version ' + str(doc.get('version_status') or 'unknown') + ' | retrieval ' + str(doc.get('status') or 'unknown'))
    lines += ['', '## 案卷轮次记录', '']
    for cycle in case.get('cycles') or []:
        lines.append('- ' + str(cycle.get('application') or '申请号未记录') +
                     ' | 最近意见 ' + str(cycle.get('latest_oa_id') or '未定位') +
                     ' | 答复候选 ' + str(cycle.get('response_id') or '未定位') +
                     ' | 配对依据 ' + str(cycle.get('pairing_status') or '未核实') +
                     ' | 目录识别数 ' + str(cycle.get('directory_count') or 0) +
                     ' | ' + str(cycle.get('source_url') or '来源链接未记录'))
    lines += ['', '最新审查意见是否已有答复尚需具体轮次证据；按日期归组的候选不等于已核实的对应关系。', '', '## 未解决的材料缺口', '']
    lines += ['- ' + gap for gap in analysis['gaps']]
    (folder / '01-同族审查历史报告.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')


def _write_summary(case, analysis, folder):
    sources = analysis['sources']
    lines = ['# 同族答复包摘要', '',
             f'目标国家：{case.get("target_country") or "BR"}。目标申请：{case.get("target_application") or "尚未填写"}。', '',
             '本包复用已选同族英文文本，生成英文答复、修改痕迹权利要求、清洁版权利要求及修改依据表。'
             '申请人的既有答辩不代表审查机关采纳；外国授权也不证明某项论证被接受。', '',
             '## 选用来源', '']
    labels = {'response': '同族答复', 'claims': '同族权利要求', 'pending': '目标在审基线', 'oa': '目标审查意见', 'original': '目标原始申请'}
    for key, label in labels.items():
        lines.append('- ' + label + '：' + _source_label(sources[key]))
    settings = case.get('settings') or {}
    lines += ['', '## 修改与审阅状态', '',
              '修改痕迹按目标英文在审基线生成。' if analysis['baseline_usable'] else '缺少可用的目标英文在审基线，修改标记页仅提供参考文本，不声称目标修改痕迹。',
              '启用了可选人工审阅。' if settings.get('review_enabled') else '未启用可选人工审阅；生成不会被该选项阻挡。']
    if settings.get('response_override') or settings.get('claims_override'):
        lines.append('导出中已标注用户自写的答复或权利要求替换文本。')
    if analysis['adaptation']['log']:
        lines.append('英文答复仅将已知同族申请号或公开号替换为目标申请号；技术论证保留原文。替换日志与原始文本均保存在来源记录中。')
    if analysis['adaptation']['foreign_legal_references']:
        lines.append('外国法律或程序引用仍待适配：' + '、'.join(analysis['adaptation']['foreign_legal_references']) + '。未自动替换为目标国家法律。')
    missing_support = list(dict.fromkeys(str(row['claim']) for row in analysis['support'] if row['status'] != 'source_located'))
    if missing_support:
        lines.append('缺少已定位原始依据的权利要求：' + '、'.join(missing_support) + '。')
    lines += ['', '期限与费用仍待核实。']
    if analysis['deadline']['nominal_date']:
        lines.append('按已选巴西初步技术要求规则算得名义日期：' + analysis['deadline']['nominal_date'] +
                     '。起算公告是否适用及节假日顺延尚未核实，不是已核定期限。')
    lines += ['', '## 待补材料和核实事项', ''] + ['- ' + gap for gap in analysis['gaps']]
    lines += ['', '英文答复保留来源的技术论证，不自动判断新颖性、创造性或新增事项。修改依据仅核对用户指定的目标原始文件页内引文，不作法律充分性结论。']
    (folder / '07-中文摘要.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')


def _write_manual_guide(case, analysis, folder):
    lines = ['# Completing the evidence package', '',
             '1. Open the target office record and identify the current substantive Office Action and immediately prior entered claims.',
             '2. Download the original files and record the application number, document role, date, language and claim version status.',
             '3. For the selected foreign member, obtain the response and claims from the same application and response cycle, or the grant claims as a reference.',
             '4. Import missing originals in the workbench. Scanned files need readable text or OCR before text based reuse.',
             '5. Select the precise evidence versions. Supply English text for an English draft; no automatic translation is performed.',
             '6. Add target original page citations for every amendment and verify each target objection, filing rule and applicable fee.',
             '7. Regenerate the package. Optional review is available but does not gate generation.', '', '## Recorded office retrieval instructions', '']
    for guide in case.get('manual_guides') or []:
        lines += ['### ' + str(guide.get('country') or 'Office') + ' ' + str(guide.get('application') or ''), '',
                  str(guide.get('url') or 'No portal URL recorded'), '', str(guide.get('instruction') or 'No office specific instruction recorded.'), '',
                  'The following is the recorded file list; completeness of the office directory remains unverified.', '']
        for document in guide.get('files') or []:
            lines.append('- ' + ' | '.join(str(document.get(key) or 'Unknown') for key in ('title', 'date', 'url', 'status')))
        if not guide.get('files'):
            lines.append('No individual file record was available from this portal.')
        lines.append('')
    lines += ['## Available source links', '']
    for document in case.get('documents') or []:
        if document.get('status') in ('manual_required', 'discovered') or not document_text(document):
            lines.append('- ' + str(document.get('title') or document.get('id') or 'Document') + ': ' +
                         str(document.get('url') or 'No document URL recorded') + '. ' + str(document.get('notes') or 'Retrieve the original from the office record.'))
    lines += ['', '## Official starting points', '',
              '- EPO Register: https://register.epo.org/',
              '- INPI patents: https://www.gov.br/inpi/pt-br/servicos/patentes',
              '- USPTO Patent Center: https://patentcenter.uspto.gov/',
              '', '## Remaining gaps', ''] + ['- ' + gap for gap in analysis['gaps']]
    (folder / '08-补充材料指南.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')


def _write_originals_bundle(case, folder, case_dir):
    """02-同族原件与目录说明.zip: every obtained family OA/response/claims
    original plus a catalog, mirroring the trial-delivery evidence folder."""
    target = str(case.get('target_country') or '').upper()
    originals = [d for d in case.get('documents', [])
                 if d.get('local_path') and d.get('status') in ('downloaded', 'imported')
                 and d.get('role') in ('oa', 'response', 'claims')
                 and str(d.get('country', '')).upper() != target]
    originals.sort(key=lambda d: (d.get('country', ''), d.get('date', ''), d.get('title', '')))
    base = Path(case_dir) if case_dir else Path(folder).resolve().parents[1]
    lines = ['# 同族审查意见与答复原件 · 目录说明', '',
             '本 ZIP 收录研究过程中实际取得的同族官方原件；本案（目标法域）材料请在案件“材料”区查看。',
             '未取得原件的同族以补件指南为准，不以公开文本冒充案卷文件。', '',
             '| 文件 | 法域 | 日期 | 角色 | 来源 | SHA256 |', '|---|---|---|---|---|---|']
    roles = {'oa': '审查意见', 'response': '申请人答复', 'claims': '权利要求'}
    members = []
    for doc in originals:
        path = (base / doc['local_path']).resolve()
        if not path.is_file() or not path.is_relative_to(base):
            continue
        name = doc.get('filename') or path.name
        members.append((path, name))
        lines.append('| {name} | {country} | {date} | {role} | {url} | {sha} |'.format(
            name=name, country=doc.get('country', ''), date=doc.get('date', ''),
            role=roles.get(doc.get('role'), doc.get('role', '')),
            url=doc.get('url', ''), sha=str(doc.get('sha256', ''))[:16]))
    if not members:
        lines.append('| （本案件尚未取得同族原件） | | | | | |')
    with zipfile.ZipFile(Path(folder) / '02-同族原件与目录说明.zip', 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('00-目录说明.md', '\n'.join(lines) + '\n')
        for path, name in members:
            archive.write(path, arcname=name)
    return len(members)


def generate_package(case, out_dir: Path, case_dir=None):
    """Write a complete draft package even when evidence is incomplete."""
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    analysis = analyze_case(case)
    proposed, before, user_authored = _claims_content(case, analysis)
    _support(analysis, proposed, before)
    if str(case.get('target_country') or '').upper() == 'US' and before:
        statuses = (analysis['sources']['pending'] or {}).get('claim_statuses') or {}
        unknown = [number for number in before if before.get(number) == proposed.get(number)
                   and (statuses.get(str(number)) or statuses.get(number)) not in ('Original', 'Previously presented', 'Withdrawn')]
        if unknown:
            analysis['gaps'].append('US historical claim statuses require verification for claims ' + ', '.join(map(str, unknown)) + '.')
        if any(number not in before and number <= max(before) for number in proposed):
            analysis['gaps'].append('A proposed new US claim number is not above every supplied prior claim number; verify the full claim numbering history.')
    analysis['gaps'] = list(dict.fromkeys(analysis['gaps']))
    _write_response(case, analysis, folder, bool(proposed))
    _write_claims(case, analysis, folder, proposed, before, user_authored)
    _write_support(analysis, folder, case.get('is_demo'))
    _write_history(case, analysis, folder)
    _write_summary(case, analysis, folder)
    _write_manual_guide(case, analysis, folder)
    if case.get('is_demo'):
        for name in ('06-修改依据表.md', '07-中文摘要.md', '08-补充材料指南.md'):
            path = folder / name
            path.write_text(DEMO_NOTICE + '\n\n' + path.read_text(encoding='utf-8'), encoding='utf-8')
    source_record = dict(is_demo=bool(case.get('is_demo')), case_notice=DEMO_NOTICE if case.get('is_demo') else '',
                         case_id=case.get('id'), input_number=case.get('input_number'),
                         target_country=case.get('target_country'), target_application=case.get('target_application'),
                         family=case.get('family') or [], documents=case.get('documents') or [],
                         events=case.get('events') or [], settings=case.get('settings') or {}, analysis=analysis,
                         method='deterministic_rules_no_llm')
    (folder / '09-来源记录.json').write_text(json.dumps(source_record, ensure_ascii=False, indent=2), encoding='utf-8')
    originals_count = _write_originals_bundle(case, folder, case_dir)
    labels = [('01-同族审查历史报告.md', '同族审查历史报告'), ('02-同族原件与目录说明.zip', f'同族原件包（{originals_count} 件+目录说明）'),
              ('03-意见陈述书.docx', '意见陈述书 Word'), ('03-意见陈述书.html', '意见陈述书预览'),
              ('04-权利要求修改标记页.docx', '权利要求修改标记页 Word'), ('04-权利要求修改标记页.html', '修改标记页预览'),
              ('05-修改后权利要求清洁稿.docx', '清洁版权利要求 Word'), ('05-修改后权利要求清洁稿.html', '清洁版权利要求预览'),
              ('06-修改依据表.md', '修改依据表'), ('06-修改依据表.csv', '修改依据表 CSV'),
              ('07-中文摘要.md', '中文摘要'), ('09-来源记录.json', '来源记录'), ('08-补充材料指南.md', '补充材料指南')]
    with zipfile.ZipFile(folder / '00-完整答复包.zip', 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, _ in labels:
            archive.write(folder / name, arcname=name)
    labels.insert(0, ('00-完整答复包.zip', '完整答复包 ZIP（含以上全部）'))
    return dict(exports=[dict(name=name, label=label) for name, label in labels], analysis=analysis, gaps=analysis['gaps'])
