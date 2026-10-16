import importlib
import json
import io
import tempfile
import unittest
import uuid
from unittest.mock import patch
from urllib.error import URLError
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]


class ClaimsTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue((ROOT / 'workbench/claims.py').exists(), 'Claim parsing is not implemented')
        self.module = importlib.import_module('workbench.claims')

    def test_claims_preserve_numbers_and_continuation(self):
        claims = self.module.parse_claims('Claims\n1. A converter comprising:\na bridge; and a controller.\n2. The converter of claim 1, wherein x > 2.')
        self.assertEqual([c['number'] for c in claims], [1, 2])
        self.assertIn('a bridge; and a controller.', claims[0]['text'])

    def test_diff_roundtrip_with_math_and_punctuation(self):
        old = 'A converter, wherein I > I1 and f = 20 kHz.'
        new = 'A converter, wherein I >= I1 and f = 25 kHz.'
        diff = self.module.diff_tokens(old, new)
        self.assertEqual(''.join(t for kind, t in diff if kind != 'insert'), old)
        self.assertEqual(''.join(t for kind, t in diff if kind != 'delete'), new)

    def test_number_normalization_and_unsafe_input(self):
        self.assertEqual(self.module.normalize_number('BR 11 2023 019754-6 A2'), 'BR112023019754A2')
        self.assertEqual(self.module.normalize_number('CN117691670A'), 'CN117691670A')
        with self.assertRaises(ValueError):
            self.module.normalize_number('../../etc/passwd')


class RetrievalTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue((ROOT / 'workbench/retrieval.py').exists(), 'Source retrieval is not implemented')
        self.module = importlib.import_module('workbench.retrieval')

    def test_challenge_is_not_pdf(self):
        with self.assertRaises(ValueError):
            self.module.validate_pdf(b'<html>Cloudflare challenge</html>')

    def test_tls_failure_uses_bounded_system_transport(self):
        with patch.object(self.module, 'build_opener') as opener, patch.object(self.module, '_curl_fetch', create=True, return_value=(b'hello', 'https://patents.google.com/', 'text/html')) as fallback:
            opener.return_value.open.side_effect = URLError('SSL: UNEXPECTED_EOF_WHILE_READING')
            data = self.module.fetch_public('https://patents.google.com/')
            self.assertEqual(data[0], b'hello')
            fallback.assert_called_once()

    def test_private_and_non_https_urls_rejected(self):
        for value in ['http://127.0.0.1/test', 'file:///secret', 'https://localhost/', 'https://register.epo.org.evil.example/']:
            with self.assertRaises(ValueError):
                self.module.validate_public_url(value)

    def test_google_family_is_taken_from_family_table(self):
        html = '''<meta name="DC.relation" content="EP21934362.1"><meta name="DC.relation" content="EP4283855A1">
        <meta name="DC.relation" content="US10574129B2"><span itemprop="title">Phase shifted converter</span>
        <table><tr itemprop="familyToFamilyCitations"><td><a href="/patent/US10574129B2/en">US10574129B2</a></td></tr></table>
        <table><tr itemprop="familyToFamily"><td itemprop="countryCode">EP</td><td><a href="/patent/EP4283855A1/en"><span itemprop="publicationNumber">EP4283855A1</span></a></td></tr></table>'''
        parsed = self.module.parse_google_patents(html, 'https://patents.google.com/patent/BR112023019754A2/en')
        self.assertIn('EP4283855A1', [r['publication'] for r in parsed['family']])
        self.assertNotIn('US10574129B2', [r['publication'] for r in parsed['family']])

    def test_ep_directory_does_not_treat_receipt_as_oa(self):
        html = '''<table><tr><td>21.11.2024</td><td><a href="javascript:NewSideWindow('application?documentId=ABC123&amp;number=EP21934362&amp;lng=en', 'doc')">European search opinion</a></td><td>4</td></tr>
        <tr><td>11.12.2025</td><td><a href="application?documentId=DEF456&amp;number=EP21934362">Electronic receipt</a></td></tr></table>'''
        docs = self.module.parse_ep_directory(html, 'EP21934362')
        self.assertEqual(len(docs), 2)
        self.assertEqual(docs[0]['role'], 'oa')
        self.assertEqual(docs[0]['date'], '2024-11-21')
        self.assertEqual(docs[1]['role'], 'other')
        self.assertEqual(docs[0]['language'], 'unknown')


class StorageTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue((ROOT / 'workbench/storage.py').exists(), 'Case storage is not implemented')
        self.module = importlib.import_module('workbench.storage')
        self.temp_path = ROOT / 'workbench-data' / 'test-runs' / uuid.uuid4().hex
        self.temp_path.mkdir(parents=True)
        self.store = self.module.Store(self.temp_path)

    def test_case_survives_reload(self):
        case = self.store.create('CN117691670A', 'BR')
        self.store.update(case['id'], lambda c: c.update(title='A test'))
        fresh = self.module.Store(self.temp_path)
        self.assertEqual(fresh.get(case['id'])['title'], 'A test')

    def test_invalid_ids_cannot_escape_storage(self):
        with self.assertRaises(ValueError):
            self.store.get('../other')

    def test_text_import_keeps_source_and_content(self):
        case = self.store.create('BR112023019754A2', 'BR')
        doc = self.store.import_document(case['id'], b'1. A converter.', {'filename':'claims.txt','role':'claims','country':'BR','version_status':'pending','language':'en','source_url':'https://example.com/source'})
        self.assertEqual(doc['text'], '1. A converter.')
        self.assertEqual(doc['url'], 'https://example.com/source')
        self.assertEqual(len(doc['sha256']), 64)

    def test_upload_name_does_not_allow_path_traversal(self):
        case = self.store.create('CN117691670A', 'BR')
        doc = self.store.import_document(case['id'], b'text', {'filename':'../../escape.txt'})
        self.assertTrue((self.store.case_dir(case['id']) / doc['local_path']).resolve().is_relative_to(self.store.case_dir(case['id'])))

    def test_marked_word_claims_are_not_silently_treated_as_clean(self):
        from docx import Document
        word = Document()
        p = word.add_paragraph('1. A converter comprising ')
        p.add_run('an old sensor').font.strike = True
        p.add_run('a controller').underline = True
        data = io.BytesIO()
        word.save(data)
        case = self.store.create('BR112023019754A2')
        doc = self.store.import_document(case['id'], data.getvalue(), {'filename':'marked.docx','role':'claims','country':'BR','version_status':'pending'})
        self.assertTrue(doc.get('requires_clean_text'))
        self.assertEqual(doc['version_status'], 'unknown')

    def test_word_table_stays_with_its_preceding_claim(self):
        from docx import Document
        from workbench.claims import parse_claims
        word = Document()
        word.add_paragraph('1. A converter having the following parameter:')
        table = word.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'alpha = 4'
        word.add_paragraph('2. The converter of claim 1 with a sensor.')
        data = io.BytesIO()
        word.save(data)
        text, pages, _ = self.module.extract_text(data.getvalue(), '.docx')
        claims = parse_claims(text)
        self.assertEqual([claim['number'] for claim in claims], [1, 2])
        self.assertIn('alpha = 4', claims[0]['text'])
        self.assertNotIn('alpha = 4', claims[1]['text'])
        self.assertEqual(pages, [{'number': None, 'text': text}])

    def test_word_nested_table_preserves_cell_content_order(self):
        from docx import Document
        from workbench.claims import parse_claims
        word = Document()
        word.add_paragraph('1. A converter with parameters:')
        table = word.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.text = 'outer before'
        nested = cell.add_table(rows=1, cols=1)
        nested.cell(0, 0).text = 'alpha = 4'
        cell.add_paragraph('outer after')
        word.add_paragraph('2. A sensor with parameters:')
        second = word.add_table(rows=1, cols=1)
        second.cell(0, 0).text = 'beta = 9'
        data = io.BytesIO()
        word.save(data)
        text, _, _ = self.module.extract_text(data.getvalue(), '.docx')
        claims = parse_claims(text)
        self.assertIn('alpha = 4', claims[0]['text'])
        self.assertLess(claims[0]['text'].index('outer before'), claims[0]['text'].index('alpha = 4'))
        self.assertLess(claims[0]['text'].index('alpha = 4'), claims[0]['text'].index('outer after'))
        self.assertNotIn('beta = 9', claims[0]['text'])
        self.assertIn('beta = 9', claims[1]['text'])


if __name__ == '__main__':
    unittest.main()
