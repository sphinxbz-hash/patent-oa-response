import copy
import importlib
import json
import shutil
import unittest
import uuid
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]


def evidence(doc_id, role, country, text, **values):
    result = dict(id=doc_id, title=doc_id, role=role, country=country, text=text,
                  date='2025-01-01', status='imported', version_status='unknown',
                  language='en', pages=[{'number': 1, 'text': text}],
                  url='https://example.org/' + doc_id, filename=doc_id + '.txt')
    result.update(values)
    return result


def example_case():
    return dict(id='example', input_number='EP1234567A1', target_country='BR',
                target_application='BR112023000001', title='Converter',
                family=[{'country': 'EP', 'application': 'EP22123456.1',
                         'publication': 'EP1234567B1', 'status': 'Reported grant',
                         'url': 'https://example.org/family', 'status_source': 'Uploaded record'}],
                events=[], settings={'source_member': 'EP', 'review_enabled': False},
                documents=[
                    evidence('ep-response', 'response', 'EP',
                             'Applicant submits that the cited document does not disclose the controller.\n'
                             'The applicant requests reconsideration.', application='EP22123456.1'),
                    evidence('ep-grant', 'claims', 'EP', '1. A converter with a controller.\n'
                             '2. The converter of claim 1 with a sensor.',
                             version_status='granted', application='EP22123456.1'),
                    evidence('br-pending', 'claims', 'BR', '1. A converter.\n'
                             '2. The converter of claim 1 with a sensor.', version_status='pending'),
                    evidence('br-oa', 'oa', 'BR', 'Inventive step\nClaim 1 lacks inventive step over EP7654321A1.',
                             pages=[{'number': 3, 'text': 'Inventive step\nClaim 1 lacks inventive step over EP7654321A1.'}]),
                    evidence('br-original', 'original', 'BR', 'The converter has a controller and a sensor.')])


class AnalysisTests(unittest.TestCase):
    def test_historical_missing_materials_clear_after_import(self):
        case = example_case()
        case['research_gaps'] = ['缺少本案可读取的审查意见', '尚未锁定本案当前有效待审权项；不会将公开文本自动升级为 pending claims']
        result = importlib.import_module('workbench.analysis').analyze_case(case)
        self.assertIsNotNone(result['sources']['oa'])
        self.assertIsNotNone(result['sources']['pending'])
        self.assertFalse(any(gap in result['gaps'] for gap in case['research_gaps']))

    def test_ep_check_digit_and_publication_selectors_resolve_same_member(self):
        case = example_case()
        case['documents'][0]['application'] = 'EP22123456'
        for member in ('EP22123456.1', 'EP22123456.1A', 'EP1234567B1'):
            case['settings']['source_member'] = member
            result = importlib.import_module('workbench.analysis').analyze_case(case)
            self.assertEqual(result['sources']['response']['id'], 'ep-response')
            self.assertEqual(result['sources']['claims']['id'], 'ep-grant')
        case['settings']['source_member'] = 'EP9999999B1'
        self.assertIsNone(importlib.import_module('workbench.analysis').analyze_case(case)['sources']['claims'])

    def setUp(self):
        self.assertTrue((ROOT / 'workbench/analysis.py').exists(), 'Case analysis is not implemented')
        self.analyze = importlib.import_module('workbench.analysis').analyze_case

    def test_source_roles_and_pending_status_are_not_inferred_from_recency(self):
        case = example_case()
        case['documents'].append(evidence('latest-br-publication', 'publication', 'BR',
                                         '1. A different device.', date='2026-10-01', version_status='publication'))
        analysis = self.analyze(case)
        self.assertEqual(analysis['sources']['pending']['id'], 'br-pending')
        self.assertEqual(analysis['sources']['claims']['id'], 'ep-grant')
        self.assertEqual(analysis['sources']['response']['id'], 'ep-response')

    def test_missing_explicit_source_does_not_silently_fall_back(self):
        case = example_case()
        case['settings']['selected_response_id'] = 'deleted-response'
        analysis = self.analyze(case)
        self.assertIsNone(analysis['sources']['response'])
        self.assertTrue(any('deleted-response' in gap for gap in analysis['gaps']))

    def test_explicit_role_and_country_mismatch_is_rejected(self):
        for key, doc_id, result_key in [('selected_response_id', 'br-oa', 'response'),
                                        ('pending_claims_id', 'ep-grant', 'pending'),
                                        ('original_id', 'ep-grant', 'original')]:
            with self.subTest(key=key):
                case = example_case()
                case['settings'][key] = doc_id
                self.assertIsNone(self.analyze(case)['sources'][result_key])

    def test_response_and_claims_do_not_cross_known_members(self):
        case = example_case()
        case['documents'].append(evidence('other-response', 'response', 'EP', 'Unrelated argument.',
                                         date='2026-10-01', application='EP99999999.9'))
        case['settings']['selected_claims_id'] = 'ep-grant'
        self.assertEqual(self.analyze(case)['sources']['response']['id'], 'ep-response')
        case['settings']['selected_response_id'] = 'other-response'
        analysis = self.analyze(case)
        self.assertIsNone(analysis['sources']['response'])
        self.assertTrue(any('member' in gap.lower() for gap in analysis['gaps']))

    def test_source_cycle_claims_without_grant_can_be_reused(self):
        case = example_case()
        case['documents'][1]['version_status'] = 'pending'
        case['documents'][1]['response_id'] = 'ep-response'
        self.assertEqual(self.analyze(case)['sources']['claims']['id'], 'ep-grant')
        case['documents'][1].pop('response_id')
        self.assertIsNone(self.analyze(case)['sources']['claims'])

    def test_latest_unpaired_response_does_not_displace_complete_response_cycle(self):
        case = example_case()
        case['documents'][1].update(version_status='pending', response_id='ep-response')
        case['documents'].append(evidence('latest-unpaired', 'response', 'EP', 'Later response without claims.',
                                         date='2026-10-01', application='EP22123456.1'))
        analysis = self.analyze(case)
        self.assertIsNotNone(analysis['sources']['claims'])
        self.assertEqual(analysis['sources']['response']['id'], 'ep-response')
        case['settings'].update(selected_claims_id='ep-grant', selected_response_id='latest-unpaired')
        self.assertIsNone(self.analyze(case)['sources']['response'])

    def test_target_evidence_with_known_different_application_is_rejected(self):
        case = example_case()
        for doc in case['documents']:
            if doc['country'] == 'BR':
                doc['application'] = 'BR112099999999'
        sources = self.analyze(case)['sources']
        self.assertIsNone(sources['pending'])
        self.assertIsNone(sources['oa'])
        self.assertIsNone(sources['original'])

    def test_response_not_entered_is_not_reused(self):
        case = example_case()
        case['documents'][0]['version_status'] = 'not_entered'
        self.assertIsNone(self.analyze(case)['sources']['response'])

    def test_manual_version_upgrade_cannot_bypass_marked_text_guard(self):
        case = example_case()
        for index in (0, 1, 2):
            case['documents'][index]['requires_clean_text'] = True
        case['documents'][1]['version_status'] = 'granted'
        case['documents'][2]['version_status'] = 'pending'
        analysis = self.analyze(case)
        for key in ('response', 'claims', 'pending'):
            self.assertIsNone(analysis['sources'][key])
        self.assertTrue(any('Marked source' in gap and 'br-pending' in gap and 'clean text required' in gap for gap in analysis['gaps']))

    def test_oa_rules_keep_page_and_patent_citation(self):
        analysis = self.analyze(example_case())
        issue = next(x for x in analysis['issues'] if x['category'] == 'inventive_step')
        self.assertEqual(issue['page'], 3)
        self.assertEqual(issue['source_id'], 'br-oa')
        self.assertIn('EP7654321A1', json.dumps(analysis['issues']))
        self.assertEqual(analysis['acceptance']['status'], 'not_verified')

    def test_empty_scan_and_missing_deadline_are_explicit_gaps(self):
        case = example_case()
        case['documents'][3].update(text='', pages=[])
        analysis = self.analyze(case)
        self.assertIsNone(analysis['sources']['oa'])
        self.assertEqual(analysis['deadline']['status'], 'unverified')
        self.assertIsNone(analysis['deadline']['due_date'])
        self.assertEqual(analysis['deadline']['fee']['status'], 'unverified')
        self.assertIsNone(analysis['deadline']['fee']['amount'])
        self.assertTrue(any('text' in gap.lower() for gap in analysis['gaps']))

    def test_br_rule_returns_nominal_date_without_claiming_verified_due_date(self):
        case = example_case()
        case['settings']['deadline'] = dict(rule_id='BR_PRELIM_90', trigger_date='2026-10-01',
                                            source_url='https://www.gov.br/inpi/pt-br/servicos/patentes/novos-guias-rapidos/guia-rapido-sobre-exigencia-tecnica-preliminar/')
        deadline = self.analyze(case)['deadline']
        self.assertEqual(deadline['nominal_date'], '2026-12-30')
        self.assertFalse(deadline['verified'])
        self.assertIsNone(deadline['due_date'])
        case['settings']['deadline']['trigger_date'] = 'not-a-date'
        self.assertIsNone(self.analyze(case)['deadline']['nominal_date'])

    def test_foreign_support_cannot_be_verified_as_target_original(self):
        case = example_case()
        case['settings']['support_rows'] = [dict(claim='1', change='controller', basis='controller', source_id='ep-grant', page=1)]
        row = self.analyze(case)['support'][0]
        self.assertEqual(row['status'], 'unverified')
        case['settings']['support_rows'][0]['source_id'] = 'br-original'
        row = self.analyze(case)['support'][0]
        self.assertEqual(row['status'], 'source_located')
        self.assertEqual(row['legal_support'], 'unverified')

    def test_strategy_candidate_distinguishes_argument_amendment_and_separate_response(self):
        case = example_case()
        case['documents'][0]['text'] += '\nEP7654321A1 describes a different arrangement.'
        analysis = self.analyze(case)
        self.assertEqual(analysis['strategy'][0]['status'], 'reuse_argument')
        self.assertEqual(analysis['strategy'][0]['method'], 'rule_based_candidate')
        case['documents'][0]['text'] = 'The applicant requests reconsideration.'
        case['documents'][3] = evidence('br-oa', 'oa', 'BR', 'Clarity of the controller is disputed.')
        self.assertEqual(self.analyze(case)['strategy'][0]['status'], 'reuse_amendment')
        case['documents'][3] = evidence('br-oa', 'oa', 'BR', 'Unity of invention is disputed.')
        self.assertEqual(self.analyze(case)['strategy'][0]['status'], 'separate_response')


class DraftingTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue((ROOT / 'workbench/drafting.py').exists(), 'Document drafting is not implemented')
        self.generate = importlib.import_module('workbench.drafting').generate_package
        self.output = ROOT / 'workbench-data' / 'test-runs' / ('drafting-' + uuid.uuid4().hex)
        self.output.mkdir(parents=True)
        self.addCleanup(self.clean_output)

    def clean_output(self):
        root = (ROOT / 'workbench-data' / 'test-runs').resolve()
        target = self.output.resolve()
        if target.is_relative_to(root) and target != root:
            shutil.rmtree(target)

    def test_complete_export_preserves_source_arguments_and_diff_formats(self):
        case = example_case()
        before = copy.deepcopy(case)
        result = self.generate(case, self.output)
        self.assertEqual(case, before)
        names = {x['name'] for x in result['exports']}
        self.assertTrue({'Response.docx', 'Response.html', 'MarkedClaims.docx', 'MarkedClaims.html',
                         'CleanClaims.docx', 'CleanClaims.html', 'SupportTable.md',
                         'SupportTable.csv', '中文摘要.md', '00-family-history.md',
                         'source-record.json', 'manual-guide.md', 'bundle.zip'} <= names)
        response = Document(self.output / 'Response.docx')
        self.assertEqual(response.paragraphs[0].style.name, 'Title')
        text = '\n'.join(p.text for p in response.paragraphs)
        self.assertIn('Applicant submits that the cited document does not disclose the controller.', text)
        self.assertIn('Applicant argument', text)
        self.assertNotIn('Chinese', text)
        marked = Document(self.output / 'MarkedClaims.docx')
        self.assertTrue(any(run.underline for p in marked.paragraphs for run in p.runs))
        clean = Document(self.output / 'CleanClaims.docx')
        self.assertFalse(any(run.underline or run.font.strike for p in clean.paragraphs for run in p.runs))
        self.assertIn('目标', (self.output / '中文摘要.md').read_text(encoding='utf-8'))
        with zipfile.ZipFile(self.output / 'bundle.zip') as package:
            self.assertNotIn('bundle.zip', package.namelist())
            self.assertIn('Response.docx', package.namelist())
        self.generate(case, self.output)
        with zipfile.ZipFile(self.output / 'bundle.zip') as package:
            self.assertNotIn('bundle.zip', package.namelist())

    def test_missing_target_baseline_never_creates_false_target_redline(self):
        case = example_case()
        case['documents'] = [d for d in case['documents'] if d['id'] != 'br-pending']
        result = self.generate(case, self.output)
        html = (self.output / 'MarkedClaims.html').read_text(encoding='utf-8')
        self.assertIn('No target redline', html)
        self.assertNotIn('<ins>', html)
        self.assertTrue(any('pending' in gap.lower() for gap in result['gaps']))

    def test_generated_missing_support_rows_propagate_to_gaps_and_summary(self):
        case = example_case()
        case['settings']['support_rows'] = [dict(claim='1', change='controller', basis='controller', source_id='br-original', page=1)]
        case['settings']['claims_override'] = '1. A converter with a controller.\n2. The converter of claim 1 with a sensor and a switch.'
        result = self.generate(case, self.output)
        self.assertTrue(any('support' in gap.lower() and '2' in gap for gap in result['gaps']))
        summary = (self.output / '中文摘要.md').read_text(encoding='utf-8')
        self.assertIn('缺少已定位原始依据的权利要求：2', summary)

    def test_missing_proposed_claims_does_not_invent_cancellations_in_support_table(self):
        case = example_case()
        case['documents'] = [d for d in case['documents'] if d['id'] != 'ep-grant']
        result = self.generate(case, self.output)
        self.assertFalse(any(row['change'] == 'Canceled claim' for row in result['analysis']['support']))
        response = (self.output / 'Response.html').read_text(encoding='utf-8')
        self.assertNotIn('Please consider the accompanying amended claims', response)
        self.assertIn('no amendment is asserted', response)

    def test_family_history_includes_reuse_argument_pages_and_evidence_completeness(self):
        case = example_case()
        paragraph = 'EP7654321A1 does not disclose the controller.'
        case['documents'][0].update(text=paragraph, pages=[{'number': 5, 'text': paragraph}])
        case['researched_at'] = '2026-10-16T01:02:03+00:00'
        self.generate(case, self.output)
        history = (self.output / '00-family-history.md').read_text(encoding='utf-8')
        self.assertIn('可复用论证清单', history)
        self.assertIn(paragraph, history)
        self.assertIn('第 5 页', history)
        self.assertIn('[来源文件](https://example.org/ep-response)', history)
        self.assertIn('Inventive step', history)
        self.assertIn('2026-10-16T01:02:03+00:00', history)
        self.assertIn('完整性未核实', history)

    def test_manual_guide_preserves_researched_portal_and_specific_file_instructions(self):
        case = example_case()
        case['manual_guides'] = [dict(country='WO', application='PCT/AA2020/123456',
            url='https://patentscope.wipo.int/search/en/search.jsf', instruction='Use Documents to retrieve the recorded written opinion.',
            files=[dict(title='Written opinion', date='2025-01-02', url='https://example.org/opinion', status='manual_required')])]
        self.generate(case, self.output)
        guide = (self.output / 'manual-guide.md').read_text(encoding='utf-8')
        self.assertIn('PCT/AA2020/123456', guide)
        self.assertIn('Use Documents to retrieve the recorded written opinion.', guide)
        self.assertIn('Written opinion | 2025-01-02', guide)
        self.assertIn('https://example.org/opinion', guide)

    def test_every_demo_artifact_visibly_identifies_fiction(self):
        case = example_case()
        case['is_demo'] = True
        result = self.generate(case, self.output)
        notice = 'Fictional demonstration / not a real patent case'
        for artifact in result['exports']:
            path = self.output / artifact['name']
            with self.subTest(name=path.name):
                if path.suffix == '.docx':
                    self.assertEqual(Document(path).paragraphs[1].text, notice)
                elif path.suffix in ('.md', '.html', '.csv', '.json'):
                    self.assertIn(notice, path.read_text(encoding='utf-8-sig'))
        record = json.loads((self.output / 'source-record.json').read_text(encoding='utf-8'))
        self.assertTrue(record.get('is_demo'))

    def test_missing_response_does_not_invent_substantive_arguments(self):
        case = example_case()
        case['documents'] = [d for d in case['documents'] if d['role'] != 'response']
        self.generate(case, self.output)
        html = (self.output / 'Response.html').read_text(encoding='utf-8')
        self.assertIn('No source backed argument', html)
        self.assertNotIn('does not disclose', html)

    def test_user_overrides_are_labelled_and_html_escaped(self):
        case = example_case()
        case['settings'].update(response_override='User argument <script>alert(1)</script>',
                                claims_override='1. A converter with a user supplied circuit.')
        self.generate(case, self.output)
        html = (self.output / 'Response.html').read_text(encoding='utf-8')
        self.assertIn('User authored', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('<script>', html)
        clean = (self.output / 'CleanClaims.html').read_text(encoding='utf-8')
        self.assertIn('user supplied circuit', clean)
        self.assertIn('User authored', clean)

    def test_duplicate_claim_numbers_are_reported_without_fabricating_claims(self):
        case = example_case()
        case['settings']['claims_override'] = '1. A device.\n1. A different device.'
        try:
            result = self.generate(case, self.output)
        except ValueError:
            self.fail('Malformed claims should create an explicit package gap, not stop generation')
        self.assertIn('No numbered proposed claims', (self.output / 'CleanClaims.html').read_text(encoding='utf-8'))
        self.assertTrue(any('parse' in gap.lower() for gap in result['gaps']))

    def test_response_adapts_exact_identity_preserves_original_and_flags_foreign_law(self):
        case = example_case()
        original = ('Application EP22123456.1 meets Article 56 EPC. '
                    'The controller differs from EP7654321A1. EP22123456.10 is another identifier.')
        case['documents'][0]['text'] = original
        result = self.generate(case, self.output)
        text = '\n'.join(p.text for p in Document(self.output / 'Response.docx').paragraphs)
        self.assertIn('Application BR112023000001 meets Article 56 EPC.', text)
        self.assertIn('EP22123456.10 is another identifier.', text)
        self.assertIn('controller differs from EP7654321A1', text)
        self.assertIn('Please consider the accompanying amended claims and following observations.', text)
        record = json.loads((self.output / 'source-record.json').read_text(encoding='utf-8'))
        self.assertEqual(record['documents'][0]['text'], original)
        self.assertTrue(result['analysis']['adaptation']['log'])
        self.assertTrue(any('EPC' in gap for gap in result['gaps']))

    def test_us_listing_has_statuses_clean_new_claims_and_canceled_number_only(self):
        case = example_case()
        case['target_country'] = 'US'
        for doc in case['documents']:
            if doc['country'] == 'BR':
                doc['country'] = 'US'
        case['settings']['claims_override'] = '1. A converter with a controller.\n3. A newly added device.'
        self.generate(case, self.output)
        doc = Document(self.output / 'MarkedClaims.docx')
        paras = [p for p in doc.paragraphs if p.text.startswith(('1.', '2.', '3.'))]
        self.assertTrue(any('1. (Currently amended)' in p.text for p in paras))
        self.assertEqual(next(p.text for p in paras if p.text.startswith('2.')), '2. (Canceled)')
        new = next(p for p in paras if p.text.startswith('3.'))
        self.assertIn('(New)', new.text)
        self.assertFalse(any(run.underline or run.font.strike for run in new.runs))


if __name__ == '__main__':
    unittest.main()
